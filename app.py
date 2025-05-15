from flask import Flask, render_template, request, redirect, url_for, flash, send_file, jsonify, session
import os
import tempfile
import shutil
import time
import threading
import uuid
from werkzeug.utils import secure_filename
import gemini  # Import your existing module
from docx import Document # Moved here as it's used in process_pdf_with_progress

app = Flask(__name__)
app.secret_key = os.urandom(24)  # Secret key for flashing messages
app.config['UPLOAD_FOLDER'] = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'uploads')
app.config['MAX_CONTENT_LENGTH'] = 64 * 1024 * 1024  # 64MB max upload size

# Create uploads folder if it doesn't exist
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

ALLOWED_EXTENSIONS = {'pdf'}
# DEFAULT_API_KEY = "" # No longer needed here, pass from frontend or config

# Dictionary to store processing tasks
processing_tasks = {}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        api_key_option = request.form.get('api_key_option', 'default')
        api_key = request.form.get('api_key') # Get custom key

        # Prioritize custom API key if provided, otherwise try to get a default (e.g., from env or config)
        # For this example, we assume if not 'default', then api_key field MUST contain the key.
        # A more robust solution would check an environment variable for the default key.
        if api_key_option == 'default':
            # Try to get default key from environment or a config file if you have one
            # For now, if default is chosen and no env var, it might fail in gemini.setup_gemini if it expects one
            # gemini.py's setup_gemini will handle None or empty api_key for now.
            api_key = os.environ.get("GEMINI_API_KEY_DEFAULT", None) 
            if not api_key:
                 # If no default key is found anywhere, and user selected default, this is an issue.
                 # However, gemini.py's setup_gemini will print an error and return None.
                 # The UI might need a way to indicate if the default key is not configured on the server.
                 print("Warning: Default API key option selected, but no default key is configured on the server.")
                 # For now, we let it proceed, gemini.py will handle the None key.

        elif not api_key: # Custom key option selected, but no key provided
            flash('Please provide your API key when selecting "Use my own API key".')
            return redirect(request.url)
        
        output_format = request.form.get('output_format')
        if not output_format:
            flash('Please select an output format')
            return redirect(request.url)
        
        if 'file' not in request.files:
            flash('No file part')
            return redirect(request.url)
        file = request.files['file']
        if file.filename == '':
            flash('No file selected')
            return redirect(request.url)
        
        if file and allowed_file(file.filename):
            original_filename = secure_filename(file.filename)
            # Create a unique internal filename to avoid conflicts if same file is uploaded multiple times
            unique_suffix = str(uuid.uuid4())[:8]
            internal_filename = f"{os.path.splitext(original_filename)[0]}_{unique_suffix}.pdf"
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], internal_filename)
            file.save(filepath)
            
            task_id = str(uuid.uuid4())
            session['current_task_id'] = task_id
            
            processing_tasks[task_id] = {
                'status': 'initializing',
                'percentage': 0,
                'current_page': 0,
                'total_pages': 0, # Will be updated by the processing function
                'message': 'Initializing processing...',
                'result': None,
                'output_format': output_format,
                'original_filename': original_filename, # Store original for user display
                'internal_filepath': filepath # Store internal path for processing
            }
            
            thread = threading.Thread(
                target=process_pdf_with_progress,
                args=(task_id, api_key, output_format)
            )
            thread.daemon = True
            thread.start()
            
            return render_template('processing.html')
        else:
            flash('File type not allowed. Please upload a PDF file.')
            return redirect(request.url)
    
    return render_template('index.html')

@app.route('/download/<filename>')
def download_file(filename):
    # Sanitize filename before joining path to prevent directory traversal
    safe_filename = secure_filename(filename) 
    if safe_filename != filename: # If secure_filename changed it, it might have been suspicious
        flash('Invalid filename.')
        return redirect(url_for('index'))

    file_path = os.path.join(app.config['UPLOAD_FOLDER'], safe_filename)
    
    if not os.path.exists(file_path):
        # Check if it's part of a task result in case the direct filename isn't found
        # This could happen if the file was renamed or if the link is from an old task
        task_found_by_file = None
        for tid, t_data in processing_tasks.items():
            if t_data.get('result'):
                if isinstance(t_data['result'], list):
                    for res_file in t_data['result']:
                        if res_file.get('name') == safe_filename:
                            task_found_by_file = t_data
                            break
                if task_found_by_file: break
        
        if not task_found_by_file:
            flash(f'The requested file "{safe_filename}" does not exist or is no longer available.')
            return redirect(url_for('index'))
    
    return send_file(file_path, as_attachment=True)

@app.errorhandler(413)
def too_large(e):
    flash('File is too large (maximum size is 64MB).')
    return redirect(url_for('index'))

def process_pdf_with_progress(task_id, api_key, output_format):
    task = processing_tasks[task_id]
    filepath = task['internal_filepath']
    original_filename_for_output = os.path.splitext(task['original_filename'])[0]
    temp_dir_from_hybrid_processing = None

    def update_task_progress(current_page, total_pages, message):
        task['current_page'] = current_page
        task['total_pages'] = total_pages
        task['message'] = message
        if total_pages > 0 and current_page > 0:
            # Rough percentage: 5% init, 90% for pages, 5% for finalization
            task['percentage'] = 5 + int((current_page / total_pages) * 90) 
        elif current_page == -1: # Special case for initial error like file not found
             task['percentage'] = 100 # End it
        else: # Initializing or early stages
            task['percentage'] = max(task['percentage'], 5) # Show some progress
        processing_tasks[task_id] = task

    try:
        task['status'] = 'processing'
        update_task_progress(0,0, 'Initializing AI model...')

        gemini_model = gemini.setup_gemini(api_key)
        if not gemini_model:
            task['status'] = 'error'
            task['message'] = 'Failed to initialize the Gemini model. Please check your API key or server configuration.'
            update_task_progress(-1, -1, task['message']) # Update with error
            return

        update_task_progress(0,0, 'Starting PDF processing...')
        
        all_extracted_text, temp_dir_from_hybrid_processing, has_errors = gemini.process_pdf_hybrid(
            filepath, 
            gemini_model, 
            update_progress_callback=update_task_progress
        )
        
        task['percentage'] = 95 # Nearing completion, before final file saving
        task['message'] = 'Generating output file...'
        processing_tasks[task_id] = task
        
        output_base = original_filename_for_output + "_extracted"
        generated_files_info = []
        
        if output_format == 'docx':
            docx_filename = output_base + ".docx"
            docx_path = os.path.join(app.config['UPLOAD_FOLDER'], docx_filename)
            document = Document()
            for idx, page_text in enumerate(all_extracted_text):
                # Add page number if desired, or keep clean
                # document.add_paragraph(f"--- Page {idx + 1} ---")
                document.add_paragraph(page_text if page_text else "[No text extracted for this page]")
                if idx < len(all_extracted_text) - 1:
                    document.add_page_break()
            document.save(docx_path)
            generated_files_info.append({'type': 'docx', 'name': docx_filename})
        
        elif output_format == 'txt':
            txt_filename = output_base + ".txt"
            txt_path = os.path.join(app.config['UPLOAD_FOLDER'], txt_filename)
            with open(txt_path, 'w', encoding='utf-8') as txt_file:
                for i, page_text in enumerate(all_extracted_text):
                    # txt_file.write(f"--- Page {i + 1} ---\\n")
                    txt_file.write(page_text if page_text else "[No text extracted for this page]")
                    if i < len(all_extracted_text) - 1:
                        txt_file.write('\n\n--- PAGE BREAK ---\n\n') # Consistent page break marker
            generated_files_info.append({'type': 'txt', 'name': txt_filename})

        task['status'] = 'completed'
        task['percentage'] = 100
        task['result'] = generated_files_info
        if has_errors:
            task['message'] = 'Processing complete with some errors. Please check the output file.'
        else:
            task['message'] = 'Processing complete! Your file is ready for download.'

    except Exception as e:
        app.logger.error(f"Error processing task {task_id}: {str(e)}", exc_info=True)
        if task_id in processing_tasks:
            task['status'] = 'error'
            task['message'] = f'An critical error occurred during processing: {str(e)}'
            task['percentage'] = 100
    finally:
        processing_tasks[task_id] = task # Ensure final task state is saved
        
        if temp_dir_from_hybrid_processing and os.path.exists(temp_dir_from_hybrid_processing):
            try:
                shutil.rmtree(temp_dir_from_hybrid_processing)
                print(f"Successfully cleaned up temp OCR image directory: {temp_dir_from_hybrid_processing}")
            except Exception as e:
                app.logger.error(f"Error cleaning up temp_dir {temp_dir_from_hybrid_processing} for task {task_id}: {str(e)}")
        
        if filepath and os.path.exists(filepath):
            try:
                os.remove(filepath)
                print(f"Successfully cleaned up uploaded PDF: {filepath}")
            except Exception as e:
                app.logger.error(f"Error cleaning up uploaded file {filepath} for task {task_id}: {str(e)}")
        
        if task_id in processing_tasks:
             processing_tasks[task_id]['original_filepath_cleaned_up'] = True

@app.route('/progress/<task_id>')
def get_progress_status(task_id):
    task = processing_tasks.get(task_id)
    if not task:
        return jsonify({'status': 'error', 'message': 'Task not found. It might have expired or never started.', 'percentage': 100}), 404
    # Ensure percentage is an int for JSON
    if 'percentage' in task and isinstance(task['percentage'], float):
        task['percentage'] = int(task['percentage'])
    return jsonify(task)

@app.route('/download-result/<task_id>') # This seems like a leftover or alternative download route?
                                      # The primary download link on processing page uses /download/<filename>
                                      # Let's ensure consistency or remove if redundant.
                                      # For now, keeping it but noting it.
def download_result_by_task_id(task_id): # Renamed to avoid conflict with download_file
    task = processing_tasks.get(task_id)
    if not task or task['status'] != 'completed' or not task.get('result'):
        flash('The requested processing task is not complete or has no results.')
        return redirect(url_for('index'))
    
    # Assuming task['result'] is a list of file_info dicts, and we download the first one for now.
    # If multiple files can be generated, this logic might need adjustment.
    if isinstance(task['result'], list) and len(task['result']) > 0:
        filename_to_download = task['result'][0].get('name')
        if filename_to_download:
            return redirect(url_for('download_file', filename=filename_to_download))
        else:
            flash('Result file name not found in task.')
            return redirect(url_for('index'))
    else:
        flash('No result files found for this task.')
        return redirect(url_for('index'))

if __name__ == '__main__':
    # For local development, consider using a default API key from environment
    # The frontend will pass the API key based on user input.
    app.run(debug=True)